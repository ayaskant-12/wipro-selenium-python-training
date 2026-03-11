
import openpyxl
from utilities.logger import get_logger

logger = get_logger()


# def get_excel_data(file_path, sheet_name):
#     """Return list of tuples (username, password) from specified sheet."""
#     try:
#         logger.info(f"Reading Excel file: {file_path}, Sheet: {sheet_name}")
#
#         workbook = openpyxl.load_workbook(file_path)
#         sheet = workbook[sheet_name]
#
#         data = []
#
#         for row in range(2, sheet.max_row + 1):
#             username = sheet.cell(row=row, column=1).value
#             password = sheet.cell(row=row, column=2).value
#
#             data.append((username, password))
#
#         logger.info(f"Total login records loaded: {len(data)}")
#
#         return data
#
#     except FileNotFoundError:
#         logger.error(f"Excel file not found: {file_path}")
#         raise
#
#     except KeyError:
#         logger.error(f"Sheet '{sheet_name}' not found in Excel file")
#         raise
#
#     except Exception as e:
#         logger.error(f"Error reading Excel file: {e}")
#         raise


def get_checkout_data(file_path, sheet_name="checkout_data"):
    """
    Read checkout details from Excel.
    Returns list of dictionaries with keys:
    first_name, last_name, address, state, postal
    """
    try:
        logger.info(f"Reading checkout data from Excel: {file_path}")

        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook[sheet_name]

        data = []

        # Read headers from first row
        headers = [cell.value for cell in sheet[1]]

        for row in range(2, sheet.max_row + 1):
            row_data = {}

            for col_idx, header in enumerate(headers, start=1):
                row_data[header] = sheet.cell(row=row, column=col_idx).value

            data.append(row_data)

        logger.info(f"Total checkout records loaded: {len(data)}")

        return data

    except FileNotFoundError:
        logger.error(f"Excel file not found: {file_path}")
        raise

    except KeyError:
        logger.error(f"Sheet '{sheet_name}' not found in Excel file")
        raise

    except Exception as e:
        logger.error(f"Error reading checkout Excel data: {e}")
        raise

